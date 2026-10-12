# -*- coding: utf-8 -*-
"""
把 judge_mcp 匯出的「轉型正義判決匯出_含威權與政治檔案條例.xlsx」轉成網站讀取的 JSON。

輸出（public/data/judgments/）：
  index.json      每件裁判的書目、主文、研究編碼（有利／否定、政治檔案／促轉個案分析）
  text/NN.json    全文分段，依序號分片，閱讀單一案件時才載入
  notes.json      各統計分頁中的文字說明（收錄標準、情境說明、法制沿革、觀察）
  persons.json    以人為單位的全案：撤銷名冊、促轉會決定書、法院裁判之間的連結與依據
  meta.json       資料來源檔名、建置時間、件數

舊版 data_source/legacy_data.js（原 public/data.js） 中 Excel 未收錄的案件（多為最高法院、最高行政法院裁定）與既有摘要會一併保留。

用法：
  python scripts/build_data.py                                  # 預設讀 ../judge_mcp/ 下的 Excel
  python scripts/build_data.py --xlsx 路徑/檔名.xlsx --legacy data_source/legacy_data.js --decisions ../tw-tj-decisions

找不到 tw-tj-decisions 時，沿用 public/data/judgments 既有的促轉會決定書與 persons.json，只更新法院裁判部分。
黨產條例的個案分析取自 data_source/party_assets_analysis.json（由 scripts/import_party_analysis.py 自 Word 整理表產生）。
"""
import argparse
import json
import os
import re
from collections import Counter, OrderedDict
from datetime import datetime
from urllib.parse import quote

import openpyxl

from link_cases import build_persons, load_decisions, load_revocations

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DEFAULT_XLSX = os.path.join(ROOT, "..", "judge_mcp", "轉型正義判決匯出_含威權與政治檔案條例.xlsx")
DEFAULT_DECISIONS = os.path.join(ROOT, "..", "tw-tj-decisions")
DEFAULT_LEGACY = os.path.join(ROOT, "data_source", "legacy_data.js")
PARTY_ANALYSIS = os.path.join(ROOT, "data_source", "party_assets_analysis.json")
FAVOR_MANUAL = os.path.join(ROOT, "data_source", "favor_manual.json")
OUT_DIR = os.path.join(ROOT, "public", "data", "judgments")

STATUTES = [
    "戒嚴時期人民受損權利回復條例", "二二八事件處理及補償條例", "戒嚴時期不當叛亂暨匪諜審判案件補償條例",
    "政黨及其附隨組織不當取得財產處理條例", "促進轉型正義條例", "威權統治時期國家不法行為被害者權利回復條例", "政治檔案條例",
    "公職人員年資併社團專職人員年資計發退離給與處理條例",
]
PARTY = "政黨及其附隨組織不當取得財產處理條例"
SENIORITY = "公職人員年資併社團專職人員年資計發退離給與處理條例"
# 條例主分頁以外、併入某條例的分頁（如黨產條例相關的最高法院民刑事判決）
EXTRA_SHEETS = {"黨產條例_最高法院": PARTY}
POS_SHEET, NEG_SHEET = "財產型_最終有利聲請人", "財產型_否定請求"
CASE_SHEETS = {"政治檔案_案件分析": "政治檔案條例", "促轉條例_案件分析": "促進轉型正義條例", "年資條例_案件分析": SENIORITY}
# 統計分頁 → 網站使用的名稱（Excel 舊名「財產型」實為二二八條例及戒嚴時期三部條例的請求案件）
NOTE_SHEETS = OrderedDict([
    ("財產型_有利類型統計", "228及戒嚴條例_有利類型統計"),
    ("財產型_否定類型統計", "228及戒嚴條例_否定類型統計"),
    ("政治檔案與促轉_類型統計", "政治檔案與促轉_類型統計"),
    ("年資條例_類型統計", "年資條例_類型統計"),
    ("年代趨勢統計", "年代趨勢統計"),
])
# 個案分析欄名統一（年資條例分頁用「上訴人（更審為原告）」等欄名）
CASE_FIELD_ALIASES = {"上訴人（更審為原告）": "原告／上訴人／聲請人", "被上訴人（更審為被告）": "被告／被上訴人／原決定機關"}
SHEET_LABEL = "228條例及戒嚴條例"
SHARDS = 24

# 法規名稱（含判決中常見簡稱）→ 正式名稱；只辨識這些法規，避免把句子片段誤判成法條
LAW_NAMES = OrderedDict([
    ("戒嚴時期人民受損權利回復條例施行細則", None), ("戒嚴時期人民受損權利回復條例", None), ("戒嚴回復條例", "戒嚴時期人民受損權利回復條例"),
    ("二二八事件處理及補償條例", None), ("二二八條例", "二二八事件處理及補償條例"),
    ("戒嚴時期不當叛亂暨匪諜審判案件補償條例", None), ("補償條例", "戒嚴時期不當叛亂暨匪諜審判案件補償條例"),
    ("政黨及其附隨組織不當取得財產處理條例", None), ("黨產條例", "政黨及其附隨組織不當取得財產處理條例"),
    ("促進轉型正義條例", None), ("促轉條例", "促進轉型正義條例"),
    ("威權統治時期國家不法行為被害者權利回復條例", None), ("威權條例", "威權統治時期國家不法行為被害者權利回復條例"),
    ("政治檔案條例", None), ("國家情報工作法", None), ("檔案法", None),
    ("懲治叛亂條例", None), ("檢肅匪諜條例", None), ("戡亂時期檢肅匪諜條例", "檢肅匪諜條例"), ("戒嚴法", None),
    ("軍事審判法", None), ("動員戡亂時期人民團體法", None), ("人民團體法", None),
    ("冤獄賠償法", None), ("刑事補償法", None), ("國家賠償法", None),
    ("行政訴訟法", None), ("民事訴訟法", None), ("刑事訴訟法", None), ("訴願法", None), ("行政程序法", None),
    ("中央法規標準法", None), ("行政罰法", None), ("中華民國刑法", "刑法"), ("刑法", None), ("民法", None),
    ("中華民國憲法", "憲法"), ("憲法增修條文", None), ("憲法", None),
    ("檢肅流氓條例", None), ("違警罰法", None), ("臺灣省戒嚴時期取締流氓辦法", None),
])
NUM = r"[0-9０-９一二三四五六七八九十百零○〇]+"
LAW_RE = re.compile(
    "(" + "|".join(re.escape(n) for n in sorted(LAW_NAMES, key=len, reverse=True)) + ")"
    r"\s*第\s*(" + NUM + r")\s*條(?:\s*之\s*(" + NUM + r"))?(?:\s*第\s*(" + NUM + r")\s*項)?"
)
INTERP_RE = re.compile(r"釋字第\s*(" + NUM + r")\s*號")
CN_DIGITS = {"零": 0, "○": 0, "〇": 0, "一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}


def to_int(s):
    s = s.translate(str.maketrans("０１２３４５６７８９", "0123456789"))
    if s.isdigit():
        return int(s)
    if all(c in CN_DIGITS for c in s):  # 四七七 這類逐位寫法
        return int("".join(str(CN_DIGITS[c]) for c in s))
    total, unit_val, num = 0, 1, 0
    for c in s:
        if c in CN_DIGITS:
            num = CN_DIGITS[c]
        elif c == "十":
            total += (num or 1) * 10; num = 0
        elif c == "百":
            total += (num or 1) * 100; num = 0
    return total + num


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def relabel(v):
    """把 Excel 說明文字中的舊分頁名稱「財產型」換成網站用語"""
    if isinstance(v, str):
        return v.replace("財產型", SHEET_LABEL)
    if isinstance(v, list):
        return [relabel(x) for x in v]
    if isinstance(v, dict):
        return {k: relabel(x) for k, x in v.items()}
    return v


def base_id(jid):
    return ",".join(str(jid).split(",")[:5])


def roc_year(date_minguo):
    m = re.search(r"(\d+)\s*年", date_minguo or "")
    return int(m.group(1)) if m else None


def iso_from_id(jid):
    parts = str(jid).split(",")
    if len(parts) >= 5 and re.fullmatch(r"\d{8}", parts[4]):
        d = parts[4]
        return f"{d[:4]}-{d[4:6]}-{d[6:]}"
    return None


def doc_type(title):
    t = title or ""
    for k in ("判決", "決定書", "裁定"):
        if t.endswith(k):
            return k
    return "其他文書"


def case_no(title):
    m = re.search(r"(\d+)\s*年度\s*(\S+?)\s*字第\s*(\d+)\s*號", title or "")
    return f"{m.group(1)}年度{m.group(2)}字第{m.group(3)}號" if m else None


def short_ref(jid):
    """TPCM,94,台覆,3,20050125 → 94台覆3（研究筆記中常用的簡稱）"""
    p = str(jid).split(",")
    return f"{p[1]}{p[2]}{p[3]}" if len(p) >= 4 else jid


def parse_sections(full):
    if not full:
        return OrderedDict()
    out = OrderedDict()
    for m in re.finditer(r"【([^】\n]{1,12})】\n(.*?)(?=\n\n【[^】\n]{1,12}】\n|\Z)", full, re.S):
        out[m.group(1)] = m.group(2).strip()
    if not out:
        out["全文"] = full.strip()
    return out


def extract_laws(sections):
    """擷取引用法條（正規化為「法規名稱第N條（第M項）」）與大法官解釋"""
    text = "\n".join(sections.values())
    seen = OrderedDict()
    for m in LAW_RE.finditer(text):
        name = LAW_NAMES[m.group(1)] or m.group(1)
        art = f"{name}第{to_int(m.group(2))}條" + (f"之{to_int(m.group(3))}" if m.group(3) else "")
        if m.group(4):
            art += f"第{to_int(m.group(4))}項"
        seen[art] = True
    for m in INTERP_RE.finditer(text):
        seen[f"釋字第{to_int(m.group(1))}號"] = True
    return list(seen)[:60]


def normalize_result(main_text, raw_result, dtype):
    """依主文細分結果；Excel 的「判決結果」只有兩類，判決類常把『原告之訴駁回』歸入撤銷。"""
    mt = main_text or ""
    if dtype == "判決":
        if "原告之訴駁回" in mt or "上訴駁回" in mt:
            return "駁回"
        if "再審之訴駁回" in mt:
            return "再審駁回"
    return raw_result or ("撤銷原判/決定" if ("撤銷" in mt or "廢棄" in mt) else "聲請駁回")


def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    header = [clean(h) for h in rows[0]]
    return header, rows[1:]


def load_legacy(path):
    if not path or not os.path.exists(path):
        return []
    src = open(path, encoding="utf-8").read()
    start, end = src.find("["), src.rfind("]")
    return json.loads(src[start:end + 1]) if start >= 0 else []


def notes_from_sheet(ws):
    """把統計分頁切成「區塊」：以【…】或「一、」開頭的列當作標題，其後的列當作表格或段落。"""
    blocks, cur = [], None
    for raw in ws.iter_rows(values_only=True):
        cells = [clean(c) for c in raw]
        while cells and cells[-1] is None:
            cells.pop()
        if not cells:
            continue
        first = cells[0] or ""
        only_one = sum(1 for c in cells if c) == 1
        is_head = only_one and (re.match(r"^[一二三四五六七八九十]+、", first) or first.startswith("【"))
        if is_head or cur is None:
            cur = {"heading": first if is_head else None, "rows": []}
            blocks.append(cur)
            if is_head:
                continue
        cur["rows"].append(cells)
    for b in blocks:
        rows = b["rows"]
        b["kind"] = "table" if rows and max(len(r) for r in rows) > 1 else "text"
        if b["kind"] == "text":
            b["paragraphs"] = [r[0] for r in rows if r and r[0]]
            del b["rows"]
    return blocks


def method_notes(rows_tail):
    """編碼分頁最下方的說明文字（收錄標準、歸類方式、更新紀錄）"""
    return [r for r in rows_tail if r]


def build(xlsx, legacy_path, decisions_dir=None):
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    legacy = {base_id(d["meta"]["id"]): d for d in load_legacy(legacy_path)}
    records = OrderedDict()   # base id -> record
    texts = {}

    # 1) 條例主分頁
    party_meta = {}   # 黨產條例分頁的編號、分類
    for sheet in STATUTES + list(EXTRA_SHEETS):
        if sheet not in wb.sheetnames:
            continue
        statute = EXTRA_SHEETS.get(sheet, sheet)
        header, rows = sheet_rows(wb[sheet])
        idx = {h: i for i, h in enumerate(header) if h}
        for r in rows:
            jid = clean(r[idx["案號"]])
            if not jid or not re.match(r"^[A-Z]{3,5},", jid):
                continue  # 底部的法院件數統計列
            b = base_id(jid)
            if "編號" in idx and clean(r[idx["編號"]]):
                party_meta[b] = {"編號": clean(r[idx["編號"]]), "分類": clean(r[idx.get("分類")]) if "分類" in idx else None}
            full = clean(r[idx["全文內容"]]) if "全文內容" in idx else None
            if b in records:
                rec = records[b]
                if statute not in rec["statutes"]:
                    rec["statutes"].append(statute)
                if full and len(full) > sum(len(v) for v in texts.get(b, {}).values()):
                    texts[b] = parse_sections(full)
                continue
            title = clean(r[idx["標題"]])
            dtype = doc_type(title)
            main = clean(r[idx["主文"]])
            records[b] = {
                "id": jid,
                "title": title,
                "caseNo": case_no(title),
                "court": clean(r[idx["法院"]]) or (title or "").split(" ")[0],
                "docType": dtype,
                "date": clean(r[idx["判決日期"]]),
                "iso": iso_from_id(jid),
                "rocYear": roc_year(clean(r[idx["判決日期"]])),
                "cause": clean(r[idx["案由"]]),
                "statute": statute,
                "statutes": [statute],
                "result": normalize_result(main, clean(r[idx["判決結果"]]), dtype),
                "mainText": main,
                "url": clean(r[idx["原始連結"]]),
                "pdf": clean(r[idx["PDF下載連結"]]),
                "source": "excel",
            }
            if full:
                texts[b] = parse_sections(full)

    # 2) 舊版資料中有、Excel 沒有的案件
    for b, d in legacy.items():
        m, am = d["meta"], d.get("analysis_meta", {})
        if b not in records:
            title = m.get("title")
            dtype = am.get("judgment_type_normalized") or doc_type(title)
            main = d.get("main_text_clean")
            records[b] = {
                "id": m["id"], "title": title, "caseNo": m.get("case_no") or case_no(title),
                "court": am.get("court_normalized") or m.get("court"), "docType": dtype,
                "date": m.get("date_minguo"), "iso": m.get("date_iso") or iso_from_id(m["id"]),
                "rocYear": roc_year(m.get("date_minguo")), "cause": m.get("cause"),
                "statute": am.get("category_normalized") or "未分類",
                "statutes": [am.get("category_normalized") or "未分類"],
                "result": normalize_result(main, d.get("decision_result"), dtype),
                "mainText": main, "url": m.get("source_url"),
                "pdf": "https://judgment.judicial.gov.tw/EXPORTFILE/ExportToPdf.aspx?type=JD&id="
                       + quote(m["id"], safe="").lower() + "&fname=" + quote(title or "", safe="") + "&lawpara=&ot=in",
                "source": "legacy",
            }
            if d.get("sections"):
                texts[b] = OrderedDict(d["sections"])
        rec = records[b]
        if d.get("ai_summary"):
            rec["summary"] = [{"point": s.get("point"), "refs": s.get("refs", [])}
                              for s in d["ai_summary"] if s.get("point")]
        if b not in texts and d.get("sections"):
            texts[b] = OrderedDict(d["sections"])

    # 3) 研究編碼：有利／否定
    coding_notes = {}
    for sheet, stance in ((POS_SHEET, "有利"), (NEG_SHEET, "否定")):
        if sheet not in wb.sheetnames:
            continue
        header, rows = sheet_rows(wb[sheet])
        idx = {h: i for i, h in enumerate(header) if h}
        tail = []
        for r in rows:
            jid = clean(r[idx["案號"]]) if "案號" in idx else None
            if not jid or not re.match(r"^[A-Z]{3,5},", jid):
                txt = clean(r[0])
                if txt:
                    tail.append(txt)
                continue
            b = base_id(jid)
            rec = records.get(b)
            if not rec:
                continue
            c = {"stance": stance}
            for h, i in idx.items():
                if h in ("案號", "標題", "法院", "裁判日期", "主文", "PDF下載連結", "原始連結", "條例"):
                    continue
                v = clean(r[i])
                if v:
                    c[h] = v
            c["codingStatute"] = clean(r[idx["條例"]]) if "條例" in idx else None
            rec["coding"] = c
            if c.get("文書類型"):
                rec["docType"] = c["文書類型"]
        coding_notes[stance] = method_notes(tail)

    # 4) 政治檔案／促轉條例個案分析
    for sheet, statute in CASE_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        header, rows = sheet_rows(wb[sheet])
        idx = {h: i for i, h in enumerate(header) if h}
        for r in rows:
            jid = clean(r[idx["案號"]])
            if not jid or not re.match(r"^[A-Z]{3,5},", jid):
                continue
            b = base_id(jid)
            rec = records.get(b)
            if not rec:
                continue
            a = {CASE_FIELD_ALIASES.get(h, h): clean(r[i]) for h, i in idx.items()
                 if h not in ("案號", "標題", "法院", "裁判日期", "PDF下載連結", "原始連結") and clean(r[i])}
            rec.setdefault("caseAnalysis", {})[statute] = a

    # 4b) 黨產條例：Word 整理表的事實簡述、上訴人、上訴結果、理由，依編號對應
    party_groups = []
    if os.path.exists(PARTY_ANALYSIS):
        pa = json.load(open(PARTY_ANALYSIS, encoding="utf-8"))
        party_groups = pa.get("groups", [])
        by_no = {row["編號"]: row for row in pa["rows"]}
        for b, pm in party_meta.items():
            row = by_no.get(pm["編號"])
            rec = records.get(b)
            if not rec:
                continue
            a = {"編號": pm["編號"], "分類": (row or {}).get("分類") or pm["分類"]}
            if row:
                a.update({k: row[k] for k in ("事實簡述", "上訴人", "上訴結果", "結果", "理由") if row.get(k)})
            rec.setdefault("caseAnalysis", {})[PARTY] = a

    # 5) 法條與文字長度
    for b, rec in records.items():
        sec = texts.get(b)
        rec["hasText"] = bool(sec)
        rec["laws"] = extract_laws(sec) if sec else []
        rec["ref"] = short_ref(rec["id"])

    # 6) 促轉會決定書併入，並以人為單位串成全案
    persons, revocations = [], []
    if not (decisions_dir and os.path.isdir(decisions_dir)):
        persons = keep_previous_decisions(records, texts)
    else:
        dec_records, dec_texts = load_decisions(decisions_dir)
        for rec in dec_records:
            rec["hasText"] = True
            rec["laws"] = extract_laws(dec_texts[rec["id"]])
            records[rec["id"]] = rec
            texts[rec["id"]] = dec_texts[rec["id"]]
        revocations = load_revocations(decisions_dir)
        persons = build_persons(list(records.values()), texts, revocations)

    sizes = {st: wb[st].max_row for st in STATUTES if st in wb.sheetnames}
    for rec in records.values():
        if len(rec["statutes"]) > 1:
            rec["statute"] = min(rec["statutes"], key=lambda st: sizes.get(st, 10 ** 6))

    ordered = sorted(records.values(), key=lambda x: (STATUTES.index(x["statute"]) if x["statute"] in STATUTES else 99,
                                                       x.get("iso") or "", x["id"]))
    for i, rec in enumerate(ordered):
        rec["shard"] = i % SHARDS if rec["hasText"] else None

    notes = {key: relabel(notes_from_sheet(wb[name])) for name, key in NOTE_SHEETS.items() if name in wb.sheetnames}
    notes["編碼說明"] = relabel(coding_notes)
    notes["黨產條例_分類"] = party_groups
    manual = json.load(open(FAVOR_MANUAL, encoding="utf-8")) if os.path.exists(FAVOR_MANUAL) else {}
    for rec in ordered:
        if rec.get("coding"):
            rec["coding"] = relabel(rec["coding"])
        rec["favor"] = favor_of(rec, manual)
    return ordered, texts, notes, persons


# ── 有利於誰 ─────────────────────────────────────────────
# 每件裁判都標示：有利於哪一方（side）、具體是誰（who）、依據（basis）、判讀方式（method），
# 以及轉型正義一方勝敗（tj）。轉型正義一方：二二八、戒嚴、威權、促轉平復等請求案件是請求人；
# 黨產、年資、政治檔案條例案件是執行轉型正義措施的國家機關。

SIDE_CLAIMANT, SIDE_STATE, SIDE_PARTY, SIDE_OFFICIAL, SIDE_PRIVATE = "請求人／被害人", "國家機關", "政黨及相關組織", "退職政務人員", "其他私人"
SIDE_MIXED, SIDE_NONE = "互有勝負", "不分勝負"
TJ_WIN, TJ_LOSE, TJ_MIXED, TJ_NA = "轉型正義一方勝", "轉型正義一方敗", "互有勝負", "不適用"
MEASURE_STATUTES = {PARTY, SENIORITY, "政治檔案條例"}
PARTY_ORGS = re.compile(r"國民黨|救國團|救助總會|自由民主聯盟|中央投資|欣裕台|欣光華|中國廣播|婦女聯合會|國家發展基金會|民權基金會|民族基金會")
STATE_ORGS = re.compile(r"支給機關|審定機關|部|局|府|委員會|院|大學|醫院|檢察署|司令部|基金會|法院|機關")
# 個別判斷：當事人身分無法由名稱判定、或轉型正義立場需另行說明者
TJ_OVERRIDE = {
    "TPBA,111,訴,260": (TJ_LOSE, "研究者申請應用政治檔案遭駁回"),
    "TPSV,109,台上,2410": (TJ_NA, "民事損害賠償"), "TPSV,112,台上,2347": (TJ_NA, "民事股東會決議爭議"),
    "TPSV,112,台上,1034": (TJ_NA, "民事損害賠償"),
}


def party_name(v):
    v = re.sub(r"^(原告|被告)：", "", v or "")
    return "、".join(x for x in (re.sub(r"[（(,，].*$", "", p).strip() for p in re.split(r"[、；]", v)) if x)


def side_of(name):
    if PARTY_ORGS.search(name or ""):
        return SIDE_PARTY
    if "政務人員" in (name or ""):
        return SIDE_OFFICIAL
    if STATE_ORGS.search(name or ""):
        return SIDE_STATE
    if re.search(r"研究者|黨史館", name or ""):
        return SIDE_PRIVATE
    return SIDE_CLAIMANT


def tj_of(side, statute, opp=None):
    """opp：敗訴一方的身分。機關勝訴而對造是政黨或政務人員時，也屬轉型正義措施獲維持"""
    if side == SIDE_MIXED:
        return TJ_MIXED
    if side in (SIDE_NONE, SIDE_PRIVATE):
        return TJ_NA
    if side == SIDE_STATE:
        return TJ_WIN if statute in MEASURE_STATUTES or opp in (SIDE_PARTY, SIDE_OFFICIAL) else TJ_LOSE
    return TJ_WIN if side == SIDE_CLAIMANT else TJ_LOSE


def favor_from_analysis(statute, a):
    """回傳 (勝訴方身分, 勝訴方, 依據, 敗訴方身分)"""
    if "結果" in a and "上訴人" in a:  # 黨產條例整理表：上訴人＋上訴結果
        app = a["上訴人"]
        app_side = side_of(app)
        basis = f"上訴人{party_name(app)}：{a.get('上訴結果') or a['結果']}"
        if a["結果"] == "發回更審":
            return app_side, party_name(app), basis, SIDE_STATE if app_side != SIDE_STATE else SIDE_PARTY
        if app_side == SIDE_STATE:
            return SIDE_PARTY, "中國國民黨", basis, SIDE_STATE
        return SIDE_STATE, "不當黨產處理委員會", basis, app_side
    out, detail = a.get("勝負歸類") or "", a.get("勝負") or ""
    first, second = a.get("原告／上訴人／聲請人", ""), a.get("被告／被上訴人／原決定機關", "")
    if "一部" in out or "雙方" in out:
        return SIDE_MIXED, f"{party_name(first)}／{party_name(second)}", detail or out, SIDE_MIXED
    # 年資條例分頁的第一欄是上訴人（常為機關），勝負歸類則以原審原告（社團）為準，故依雙方身分判斷
    state, other = (first, second) if side_of(first) == SIDE_STATE else (second, first)
    if side_of(first) == SIDE_STATE or side_of(second) == SIDE_STATE:
        claimant_won = out in ("請求人勝訴（含廢棄發回）", "原告勝訴（含廢棄發回）")
        winner, loser = (other, state) if claimant_won else (state, other)
        return side_of(winner), party_name(winner), out, side_of(loser)
    winner, loser = (first, second) if "勝訴（含" in out else (second, first)
    return side_of(winner), party_name(winner), out, side_of(loser)


def favor_of(rec, manual):
    st = rec["statute"]
    m = manual.get(rec["id"]) or manual.get(base_id(rec["id"]))
    if m:
        f = dict(m, method="主文與當事人判讀")
    elif rec.get("caseAnalysis"):
        sts = [x for x in rec["caseAnalysis"] if x in MEASURE_STATUTES] or list(rec["caseAnalysis"])
        st = sts[0]
        side, who, basis, opp = favor_from_analysis(st, rec["caseAnalysis"][st])
        f = {"side": side, "who": who, "basis": basis, "method": "個案分析", "_opp": opp}
    elif rec.get("coding"):
        c = rec["coding"]
        if c["stance"] == "有利":
            f = {"side": SIDE_CLAIMANT, "who": "聲請人／原告", "basis": c.get("有利結果類型") or "有利於聲請人", "method": "Excel 歸類"}
        else:
            f = {"side": SIDE_STATE, "who": "國家（請求遭否定）", "basis": c.get("否定類型") or "否定請求", "method": "Excel 歸類"}
    elif rec.get("source") == "tjc":
        res = rec.get("result") or ""
        if res.startswith("平復"):
            f = {"side": SIDE_CLAIMANT, "who": "、".join(rec.get("subjects") or []) or "聲請人", "basis": res, "method": "決定書結果"}
        elif res == "駁回":
            f = {"side": SIDE_STATE, "who": "促進轉型正義委員會", "basis": "平復聲請駁回", "method": "決定書結果"}
        else:
            f = {"side": SIDE_NONE, "who": "—", "basis": "調查報告，非准駁決定", "method": "決定書結果"}
    else:
        return None
    f["tj"] = tj_of(f["side"], st, f.pop("_opp", None))
    ov = TJ_OVERRIDE.get(base_id(rec["id"])) or TJ_OVERRIDE.get(",".join(rec["id"].split(",")[:4]))
    if ov:
        f["tj"], f["tjNote"] = ov
    return f


def keep_previous_decisions(records, texts):
    """沒有 tw-tj-decisions 時，從上次的輸出取回促轉會決定書與全案連結"""
    idx_path = os.path.join(OUT_DIR, "index.json")
    if not os.path.exists(idx_path):
        return []
    prev = [r for r in json.load(open(idx_path, encoding="utf-8")) if r.get("source") == "tjc"]
    shards = {}
    for rec in prev:
        if rec.get("shard") is not None and rec["shard"] not in shards:
            shards[rec["shard"]] = json.load(open(os.path.join(OUT_DIR, "text", f"{rec['shard']:02d}.json"), encoding="utf-8"))
        records[rec["id"]] = rec
        texts[rec["id"]] = shards.get(rec.get("shard"), {}).get(rec["id"], {})
        rec["hasText"] = bool(texts[rec["id"]])
    p_path = os.path.join(OUT_DIR, "persons.json")
    return json.load(open(p_path, encoding="utf-8")) if os.path.exists(p_path) else []


def write(ordered, texts, notes, xlsx, persons=()):
    os.makedirs(os.path.join(OUT_DIR, "text"), exist_ok=True)
    shards = [dict() for _ in range(SHARDS)]
    for rec in ordered:
        if rec["shard"] is not None:
            shards[rec["shard"]][rec["id"]] = texts[base_id(rec["id"])]
    for i, s in enumerate(shards):
        with open(os.path.join(OUT_DIR, "text", f"{i:02d}.json"), "w", encoding="utf-8") as f:
            json.dump(s, f, ensure_ascii=False, separators=(",", ":"))
    with open(os.path.join(OUT_DIR, "index.json"), "w", encoding="utf-8") as f:
        json.dump(ordered, f, ensure_ascii=False, separators=(",", ":"))
    with open(os.path.join(OUT_DIR, "notes.json"), "w", encoding="utf-8") as f:
        json.dump(notes, f, ensure_ascii=False, separators=(",", ":"))
    with open(os.path.join(OUT_DIR, "persons.json"), "w", encoding="utf-8") as f:
        json.dump(persons, f, ensure_ascii=False, separators=(",", ":"))
    meta = {
        "source": os.path.basename(xlsx),
        "sourceModified": datetime.fromtimestamp(os.path.getmtime(xlsx)).isoformat(timespec="minutes"),
        "builtAt": datetime.now().isoformat(timespec="minutes"),
        "total": len(ordered),
        "withText": sum(1 for r in ordered if r["hasText"]),
        "coded": dict(Counter(r["coding"]["stance"] for r in ordered if r.get("coding"))),
        "favor": dict(Counter(r["favor"]["side"] for r in ordered if r.get("favor"))),
        "unclassified": sum(1 for r in ordered if not r.get("favor")),
        "caseAnalysis": sum(1 for r in ordered if r.get("caseAnalysis")),
        "byStatute": dict(Counter(r["statute"] for r in ordered)),
        "bySource": dict(Counter(r["source"] for r in ordered)),
        "persons": len(persons),
        "personsLinked": sum(1 for p in persons if len(p["kinds"]) > 1),
        "personsWithCourt": sum(1 for p in persons if p["court"]),
    }
    with open(os.path.join(OUT_DIR, "meta.json"), "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)
    return meta


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--legacy", default=DEFAULT_LEGACY)
    ap.add_argument("--decisions", default=DEFAULT_DECISIONS, help="tw-tj-decisions 專案目錄（促轉會決定書與撤銷名冊）")
    args = ap.parse_args()
    ordered, texts, notes, persons = build(args.xlsx, args.legacy, args.decisions)
    print(json.dumps(write(ordered, texts, notes, args.xlsx, persons), ensure_ascii=False, indent=2))
